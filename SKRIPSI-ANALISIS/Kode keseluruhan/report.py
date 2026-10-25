"""Tabel perbandingan & ekspor laporan Excel (Pandas + Openpyxl) - Bab 4.2.6 poin 7."""
import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import database as db

JENIS_LAPORAN = ["Laporan Lengkap", "Data Transaksi", "Hasil Apriori", "Hasil FP-Growth",
                 "Perbandingan Algoritma", "Rekomendasi Bisnis"]


def pct(x):
    return f"{x * 100:.2f}%"


def itemset_df(items):
    return pd.DataFrame([{"No": i, "Itemset": ", ".join(d["itemset"]), "Support": pct(d["support"]),
                          "Jumlah Trx": d["count"]} for i, d in enumerate(items, 1)])


def rules_df(rules):
    return pd.DataFrame([{"No": i, "Antecedent": ", ".join(r["antecedent"]), "Consequent": ", ".join(r["consequent"]),
                          "Support": pct(r["support"]), "Confidence": pct(r["confidence"]), "Lift": round(r["lift"], 3)}
                         for i, r in enumerate(rules, 1)])


def comparison_df(analyses: dict) -> pd.DataFrame:
    rows = []
    for name in ("Apriori", "FP-Growth"):
        a = analyses.get(name)
        if not a:
            continue
        rl = a["rules"]
        rows.append({
            "Algoritma": name, "Total Frequent": a["jumlah_frequent_itemset"], "Total Aturan": a["jumlah_aturan_asosiasi"],
            "Waktu Eksekusi (detik)": round(a["waktu_eksekusi"], 4),
            "Support Tertinggi": pct(max(r["support"] for r in rl)) if rl else "-",
            "Confidence Tertinggi": pct(max(r["confidence"] for r in rl)) if rl else "-",
            "Lift Tertinggi": round(max(r["lift"] for r in rl), 3) if rl else "-"})
    return pd.DataFrame(rows)


def fp_frames(fp: dict):
    """Header table & CPB dalam bentuk tabel (untuk layar maupun Excel)."""
    header = pd.DataFrame([{"No": i, "Item": h["item"], "Jumlah Trx": h["count"], "Support": pct(h["support"]),
                            "Status": "Frequent"} for i, h in enumerate(fp["header"], 1)])
    rows = []
    for item, d in fp["cpb"].items():
        for i, p in enumerate(d["paths"], 1):
            rows.append({"Item": item, "No": i, "Prefix path": " → ".join(p["path"]), "Count node": p["count"]})
    return header, pd.DataFrame(rows)


def _style(writer):
    for ws in writer.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="F1EEE6")
            c.fill = PatternFill("solid", fgColor="22373A")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, col in enumerate(ws.columns, 1):
            w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(i)].width = min(max(w + 3, 10), 70)
        ws.freeze_panes = "A2"


def build_excel(jenis: str, start=None, end=None):
    """Kembalikan (bytes, daftar_peringatan)."""
    warn, sheets = [], {}
    need_result = jenis != "Data Transaksi"
    code = db.latest_code()
    an = db.get_analysis_by_code(code) if code else {}
    if need_result and not an:
        warn.append("Belum ada hasil analisis. Jalankan analisis terlebih dahulu.")

    if jenis in ("Laporan Lengkap", "Data Transaksi"):
        sheets["Data Transaksi"] = db.get_transactions(start, end).rename(columns={
            "no_transaksi": "No_Transaksi", "tanggal": "Tanggal", "nama_produk": "Nama_Produk",
            "jumlah": "Jumlah", "harga_satuan": "Harga_Satuan", "total": "Total"})
    if an:
        algos = {"Laporan Lengkap": ["Apriori", "FP-Growth"], "Hasil Apriori": ["Apriori"],
                 "Hasil FP-Growth": ["FP-Growth"]}.get(jenis, [])
        if jenis in ("Laporan Lengkap", "Hasil Apriori", "Hasil FP-Growth"):
            sheets["Parameter Analisis"] = pd.DataFrame([{
                "Kode": a["comparison_code"], "Algoritma": a["nama_algoritma"], "Min Support (%)": a["minimum_support"],
                "Min Confidence (%)": a["minimum_confidence"], "Min Lift": f'>{a["minimum_lift"]:g}',
                "Jumlah Frequent": a["jumlah_frequent_itemset"], "Jumlah Aturan": a["jumlah_aturan_asosiasi"],
                "Waktu Eksekusi (detik)": round(a["waktu_eksekusi"], 4), "Tanggal Analisis": a["tanggal_analisis"]}
                for k, a in an.items() if k in algos])
        for k in algos:
            if k in an:
                tag = "Apriori" if k == "Apriori" else "FP-Growth"
                sheets[f"Frequent {tag}"] = itemset_df(an[k]["freq"]["itemsets"])
                sheets[f"Aturan {tag}"] = rules_df(an[k]["rules"])
        if "FP-Growth" in algos and "FP-Growth" in an and an["FP-Growth"]["freq"].get("fp"):
            h, c = fp_frames(an["FP-Growth"]["freq"]["fp"])
            sheets["FP-Tree Header Table"], sheets["FP-Growth CPB"] = h, c
        if jenis in ("Laporan Lengkap", "Perbandingan Algoritma"):
            sheets["Perbandingan"] = comparison_df(an)
        if jenis in ("Laporan Lengkap", "Rekomendasi Bisnis"):
            rec = db.get_results(code, "Rekomendasi")
            sheets["Rekomendasi Bisnis"] = pd.DataFrame(
                [{"No": i, "Rekomendasi": r["judul"], "Penjelasan": r["deskripsi"]} for i, r in enumerate(rec, 1)])
    if not sheets:
        return None, warn
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, df in sheets.items():
            df.to_excel(w, sheet_name=name[:31], index=False)
        _style(w)
    return buf.getvalue(), warn


def file_name(jenis):
    return f"{jenis.lower().replace(' ', '_')}_{datetime.now():%Y%m%d_%H%M}.xlsx"
